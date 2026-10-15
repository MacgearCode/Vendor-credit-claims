"""Sell-Through Table — the weekly team workbook as a filterable web table, with
an Excel download of exactly what is on screen.

Same measures and the same SQL shape as `Sell Through/db/export_matrix.py`
(itself verified cell-for-cell against reporting.v_sell_through_customer_week):

    Retail Value = Qty x RRP                         (anchor week)
    6Wk Avg      = sum(Qty over the trailing 6 weeks) / 6
    WOS          = SOH / 6Wk Avg, clamped to 0..99; SOH <= 0 -> 0; no sales -> 99
    then Qty for each prior week, newest first.

Change the definitions in both places together.

Grains:
    retailer  product x retailer, product-major — the workbook's "All AU" tab
    product   one row per product, retailers rolled up
    store     product x store for ONE retailer — the workbook's "<Retailer> Stores" tabs

Scope (the user's brands and regions) is applied in SQL like every report.
"""
from datetime import date, timedelta
from io import BytesIO

from . import warehouse

WOS_LOW, WOS_HIGH = 3.0, 12.0          # same bands as the workbook
LOOKBACK_DAYS = 35                     # the 6-week window reaches back 35 days
GRAINS = ("retailer", "product", "store")
MAX_PRIOR_WEEKS = 52


def options(brands, region: str) -> dict:
    """Filter choices in scope: weeks loaded, retailers and brands with data.
    Reads the fact on its own keys (region_key, date_key) and names them after:
    ~0.4 s, against ~3 s grouping through the dimension joins."""
    def build():
        weeks = [w.isoformat() for w in sorted(warehouse.recent_weeks(104), reverse=True)]
        since = int((date.today() - timedelta(days=400)).strftime("%Y%m%d"))
        rows = warehouse.query("""
            with f as (
                select f.customer_key, f.item_key, bool_or(f.store_name is not null) as stores
                from reporting.fact_sell_through f
                where f.region_key = %(rk)s and f.date_key > %(since)s
                group by 1, 2)
            select c.customer_name, bool_or(f.stores), i.brand
            from f
            join reporting.dim_customer c on c.customer_key = f.customer_key
            join reporting.dim_item     i on i.item_key     = f.item_key
            where (%(brands)s::text[] is null or i.brand = any(%(brands)s::text[]))
            group by 1, 3""", {"rk": warehouse.region_key(region), "since": since, "brands": brands})
        retailers, has_stores, brand_set = set(), set(), set()
        for cust, stores, brand in rows:
            retailers.add(cust)
            if stores:
                has_stores.add(cust)
            if brand:
                brand_set.add(brand)
        return {"weeks": weeks, "retailers": sorted(retailers),
                "store_retailers": sorted(has_stores), "brands": sorted(brand_set)}
    key = ("table_opts", None if brands is None else tuple(sorted(brands)), region)
    return warehouse.cached(key, build, ttl=warehouse.SELL_THROUGH_TTL)


def _keys(grain):
    if grain == "product":
        return ["code", "description", "brand"], "item_key"
    if grain == "store":
        return (["code", "description", "brand", "retailer", "store"],
                "item_key, retailer, store")
    return (["code", "description", "brand", "retailer", "customer_code"],
            "item_key, retailer")


def build(brands, region: str, grain: str = "retailer", anchor: date | None = None,
          prior: int = 12, brand: str = "", retailer: str = "") -> dict:
    grain = grain if grain in GRAINS else "retailer"
    prior = max(0, min(int(prior), MAX_PRIOR_WEEKS))
    if grain == "store" and not retailer:
        return {"error": "Pick a retailer to see its stores.", "rows": [], "weeks": []}
    if brands == []:
        return {"rows": [], "weeks": [], "anchor": None}

    # Aggregate the fact on its own indexed keys first (region_key, date_key range),
    # then attach names: the same rows as grouping through the dimension joins,
    # in about half the time, and only the weeks shown come back over the wire.
    where = ["f.region_key = %(rk)s", "f.date_key between %(lo)s and %(hi)s"]
    p = {"rk": warehouse.region_key(region), "brands": brands}
    if brands is not None or brand:
        where.append("""f.item_key in (select item_key from reporting.dim_item
                          where (%(brands)s::text[] is null or brand = any(%(brands)s::text[]))
                            and (%(brand)s = '' or brand = %(brand)s))""")
        p["brand"] = brand
    if retailer:
        where.append("f.customer_key in (select customer_key from reporting.dim_customer"
                     " where customer_name = %(retailer)s)")
        p["retailer"] = retailer
    if grain == "store":
        where.append("f.store_name is not null")
    w = " and ".join(where)

    # Week columns come from the global list of loaded weeks (every retailer
    # loads the same weeks): consistent columns, and ~0.1 s instead of ~10 s
    # through the joins. A filter with no sales that week shows a blank column.
    week_list = sorted(warehouse.recent_weeks(prior + 1, anchor), reverse=True)
    if not week_list:
        return {"rows": [], "weeks": [], "anchor": None}
    if anchor is not None and week_list[0] != anchor:
        return {"rows": [], "weeks": [], "anchor": anchor.isoformat(),
                "error": "No data loaded for that week."}
    anchor = week_list[0]
    lo = min(week_list) - timedelta(days=LOOKBACK_DAYS)

    fsel = {"product": "", "retailer": ", f.customer_key", "store": ", f.customer_key, f.store_name"}[grain]
    code_agg = (", string_agg(distinct f.customer_product_code, ' / ') as customer_code"
                if grain == "retailer" else "")
    keys, _ = _keys(grain)
    part = {"product": "item_key", "retailer": "item_key, customer_key",
            "store": "item_key, customer_key, store_name"}[grain]
    out_cols = {"code": "i.macgear_code", "description": "i.description", "brand": "i.brand",
                "retailer": "c.customer_name", "customer_code": "a.customer_code",
                "store": "a.store_name"}
    join_c = ("join reporting.dim_customer c on c.customer_key = a.customer_key"
              if grain != "product" else "")
    with warehouse.connect() as conn, conn.cursor() as cur:
        cur.execute(f"""
            with f as (
                select f.item_key{fsel}, f.date_key{code_agg},
                       sum(f.qty_sold) as qty, sum(f.soh) as soh,
                       sum(f.qty_sold * f.rrp) as retail_value
                from reporting.fact_sell_through f
                where {w}
                group by f.item_key{fsel}, f.date_key),
            a as (
                select *, to_date(date_key::text, 'YYYYMMDD') as wk,
                       round(sum(qty) over win / 6.0, 2) as avg6,
                       case when soh is null then null
                            when soh <= 0 then 0
                            when sum(qty) over win / 6.0 > 0
                                then round(least(soh / (sum(qty) over win / 6.0), 99), 2)
                            else 99 end as wos
                from (select *, to_date(date_key::text, 'YYYYMMDD') as d from f) f
                window win as (partition by {part} order by d
                               range between interval '35 days' preceding and current row))
            select {', '.join(out_cols[k] for k in keys)}, a.wk, a.qty, a.soh,
                   round(a.retail_value, 2), a.avg6, a.wos
            from a
            join reporting.dim_item i on i.item_key = a.item_key
            {join_c}
            where a.wk = any(%(weeks)s)""",
            {**p, "lo": int(lo.strftime("%Y%m%d")), "hi": int(anchor.strftime("%Y%m%d")),
             "weeks": week_list})
        raw = cur.fetchall()

    nk = len(keys)
    widx = {wk: k for k, wk in enumerate(week_list)}
    rows: dict = {}
    for r in raw:
        key = tuple(r[:nk])
        wk, qty, soh, rv, avg6, wos = r[nk:]
        if wk not in widx:
            continue
        row = rows.setdefault(key, {"k": list(key), "rv": None, "avg6": None, "wos": None,
                                    "soh": None, "q": [None] * len(week_list)})
        k = widx[wk]
        row["q"][k] = None if qty is None else int(qty)
        if k == 0:
            row["rv"] = None if rv is None else float(rv)
            row["avg6"] = None if avg6 is None else float(avg6)
            row["wos"] = None if wos is None else float(wos)
            row["soh"] = None if soh is None else int(soh)

    out = list(rows.values())
    # product-major like the workbook: products by total retail value, then lines within
    tot: dict = {}
    for row in out:
        tot[row["k"][0]] = tot.get(row["k"][0], 0) + (row["rv"] or 0)
    out.sort(key=lambda row: (-tot[row["k"][0]], row["k"][0] or "", -(row["rv"] or 0),
                              [str(x or "") for x in row["k"][1:]]))
    return {"grain": grain, "keys": keys, "anchor": anchor.isoformat(),
            "weeks": [wk.isoformat() for wk in week_list], "rows": out,
            "region": region, "currency": "NZD" if region == "NZ" else "AUD"}


def build_for(brands, region, grain, anchor, prior, brand, retailer):
    key = ("table", None if brands is None else tuple(sorted(brands)), region, grain,
           anchor, prior, brand, retailer)
    return warehouse.cached(key, lambda: build(brands, region, grain, anchor, prior, brand, retailer),
                            ttl=warehouse.SELL_THROUGH_TTL)


# ---- Excel -------------------------------------------------------------------
HEADERS = {"code": "Code", "description": "Description", "brand": "Brand",
           "retailer": "Customer Name", "customer_code": "Customer Code", "store": "Store"}
WIDTHS = {"Code": 16, "Description": 46, "Brand": 15, "Customer Name": 18,
          "Customer Code": 15, "Store": 30, "Retail Value": 13, "6Wk Avg": 10,
          "WOS": 8, "SOH": 9, "Qty": 8}
NAVY = "1F4E79"


def week_label(iso: str) -> str:
    d = date.fromisoformat(iso)
    return f"Wk {d.isocalendar().week}\n{d.strftime('%d/%m/%Y')}"


def to_xlsx(t: dict, title: str, filters: str) -> bytes:
    """The table as the workbook lays it out: key columns, the anchor week's
    measure block, then Qty by prior week. Values, not formulas."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = f"{t.get('region', '')} {t.get('grain', '')}".strip()[:31] or "Sell-Through"
    hdr = [HEADERS[k] for k in t.get("keys", [])] + ["Retail Value", "6Wk Avg", "WOS", "SOH", "Qty"]
    hdr += [week_label(wk) for wk in t.get("weeks", [])[1:]]
    ws.cell(row=1, column=1, value=f"{title} — {filters}").font = Font(bold=True, size=11)
    fill = PatternFill("solid", fgColor=NAVY)
    for j, h in enumerate(hdr, start=1):
        c = ws.cell(row=2, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS.get(h, 11)
    ws.row_dimensions[2].height = 30

    nk = len(t.get("keys", []))
    money = "$#,##0.00" if t.get("currency") != "NZD" else '"NZ$"#,##0.00'
    thin = Side(style="thin", color="D9D9D9")
    for i, row in enumerate(t.get("rows", []), start=3):
        vals = row["k"] + [row["rv"], row["avg6"], row["wos"], row["soh"], row["q"][0]] + row["q"][1:]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=j, value=v)
            if j == nk + 1:
                c.number_format = money
            elif j in (nk + 2, nk + 3):
                c.number_format = "0.00"
            elif j > nk + 3:
                c.number_format = "#,##0"
            if j > nk + 5:
                c.font = Font(color="808080")
                c.border = Border(left=thin)
        wc = ws.cell(row=i, column=nk + 3)
        if isinstance(wc.value, (int, float)):
            wc.font = Font(color="C00000", bold=True) if wc.value < WOS_LOW else (
                Font(color="BF8F00") if wc.value > WOS_HIGH else Font(color="2E7D32"))
    ws.freeze_panes = ws.cell(row=3, column=nk + 6)
    last = len(t.get("rows", [])) + 2
    ws.auto_filter.ref = f"A2:{get_column_letter(len(hdr))}{max(last, 2)}"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
